"""Render the week-by-week plan, the paper list and the Excel tracker.

    python tools/build.py

Outputs:
    curriculum/04-week-by-week.md
    curriculum/07-paper-reading-list.md
    tracker/AI_Engineer_2026_Plan.xlsx
"""

import os
import sys
from collections import OrderedDict
from datetime import timedelta

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.dirname(__file__))
from plan_data import (END, KNOWN, LEARN, PAPER_PREREQS, PAPERS, PHASES,  # noqa: E402
                       PROJECTS, REPO, SATURDAY_HOURS, START, U, WEEKDAY_HOURS)
from plan_weeks import WEEKS  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
GITHUB_BLOB = "https://github.com/anshu8830/anshu/blob/main/"
DAY_NAMES = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]
PAPER_BY_KEY = {p[0]: p for p in PAPERS}


# ---------------------------------------------------------------- helpers
def resolve(link):
    """Return (url_for_excel, markdown_link_target) for a link spec."""
    if not link:
        return "", ""
    if link.startswith("repo:"):
        path = REPO[link[5:]]
        return GITHUB_BLOB + path, os.path.relpath(path, "curriculum")
    if link.startswith("http"):
        return link, link
    url = U[link]
    return url, url


def week_dates(w):
    start = START + timedelta(days=7 * (w["n"] - 1))
    return start, start + timedelta(days=len(w["days"]) - 1)


def fmt(d):
    return d.strftime("%a %d %b")


def rows():
    """Flatten the plan into one row per task."""
    out = []
    for w in WEEKS:
        start, _ = week_dates(w)
        for i, day in enumerate(w["days"]):
            d = start + timedelta(days=i)
            for t in day:
                phase, typ, topic, desc, link, hours, deliv = t
                out.append(dict(week=w["n"], date=d, day=DAY_NAMES[d.weekday()], phase=phase, type=typ,
                                topic=topic, desc=desc, link=link, hours=hours, deliv=deliv))
    return out


def validate():
    problems = []
    for w in WEEKS:
        start, _ = week_dates(w)
        for i, day in enumerate(w["days"]):
            d = start + timedelta(days=i)
            if d > END:
                problems.append(f"W{w['n']} day {i} is after END")
            total = sum(t[5] for t in day)
            expected = SATURDAY_HOURS if d.weekday() == 5 else WEEKDAY_HOURS
            if total not in (expected, 0.0):
                problems.append(f"W{w['n']} {d}: {total} h (expected {expected})")
            for t in day:
                if t[0] not in PHASES:
                    problems.append(f"unknown phase {t[0]}")
                try:
                    resolve(t[4])
                except KeyError as e:
                    problems.append(f"W{w['n']} unknown link key {e}")
        for key in w["papers"]:
            if key not in PAPER_BY_KEY:
                problems.append(f"W{w['n']} unknown paper {key}")
        for group in ("primary", "secondary", "reference"):
            for _, link in w[group]:
                try:
                    resolve(link)
                except KeyError as e:
                    problems.append(f"W{w['n']} {group}: unknown link key {e}")
        for v in w["videos"]:
            try:
                resolve(v[2])
            except KeyError as e:
                problems.append(f"W{w['n']} video: unknown link key {e}")
    for p in PAPERS:
        resolve(p[0])
    return problems


def esc(text):
    """Escape characters that would break a markdown table cell."""
    return str(text).replace("|", "\\|")


def md_link(label, link):
    _, target = resolve(link)
    return f"[{esc(label)}]({target})" if target else esc(label)


# ---------------------------------------------------------------- markdown
def render_weeks_md():
    all_rows = rows()
    total_hours = sum(r["hours"] for r in all_rows)
    lines = [
        "# 4. Week-by-week curriculum (28 Sep - 31 Dec 2026)",
        "",
        "> Generated from `tools/plan_data.py` + `tools/plan_weeks.py` by `python tools/build.py`. "
        "Edit the data files, not this page.",
        "",
        f"**{len(WEEKS)} weeks - {len({r['date'] for r in all_rows})} study days - {total_hours:.0f} planned hours** "
        f"(Mon-Fri {WEEKDAY_HOURS} h, Sat {SATURDAY_HOURS} h, Sunday = rest/catch-up; 25 Dec is a buffer day).",
        "",
        "How to read each week: **PRIMARY** is what you actually follow; **SECONDARY** is an optional alternative "
        "explanation if the primary does not click; **REFERENCE** is for lookup only. Every paper is read with the "
        "three-pass method and gets a half-page note in your repo (`papers/NN-name.md`).",
        "",
        "| Week | Dates | Focus | Phases | Hours | Project |",
        "|---|---|---|---|---|---|",
    ]
    for w in WEEKS:
        s, e = week_dates(w)
        hrs = sum(t[5] for d in w["days"] for t in d)
        lines.append(f"| [W{w['n']}](#week-{w['n']}) | {fmt(s)} - {fmt(e)} | {esc(w['title'])} | {esc(w['phases'])} | {hrs:g} | {esc(w['project'])} |")
    lines.append("")

    for w in WEEKS:
        s, e = week_dates(w)
        hrs = sum(t[5] for d in w["days"] for t in d)
        lines += [
            "---", "",
            f'<a id="week-{w["n"]}"></a>',
            f"## Week {w['n']} - {w['title']}",
            "",
            f"**Dates:** {fmt(s)} - {fmt(e)} 2026 - **Phases:** {w['phases']} - **Planned hours:** {hrs:g}",
            "",
            f"**Goal.** {w['goal']}",
            "",
            f"**Learning flow.** {w['flow']}",
            "",
        ]
        if w["primary"] or w["secondary"] or w["reference"]:
            lines.append("### Reading")
            for tier, key in (("PRIMARY", "primary"), ("SECONDARY", "secondary"), ("REFERENCE", "reference")):
                if w[key]:
                    items = "; ".join(md_link(lbl, lk) for lbl, lk in w[key])
                    lines.append(f"- **{tier}:** {items}")
            lines.append("")
        if w["videos"]:
            lines += ["### Videos", "", "| Video | Instructor / channel | Duration | Why watch | Covers |", "|---|---|---|---|---|"]
            for title, chan, link, dur, why, sec in w["videos"]:
                lines.append(f"| {md_link(title, link)} | {esc(chan)} | {dur} | {esc(why)} | {esc(sec)} |")
            lines.append("")
        if w["papers"]:
            lines += ["### Papers", "", "| Paper | Authors (year) | Read carefully | Implement |", "|---|---|---|---|"]
            for key in w["papers"]:
                p = PAPER_BY_KEY[key]
                lines.append(f"| {md_link(p[1], key)} | {esc(p[2])} ({p[3]}) | {esc(p[7])} | {esc(p[8])} |")
            lines.append("")
        lines += ["### Coding exercises", ""] + [f"- {c}" for c in w["coding"]] + [""]
        lines += [f"### Project: {w['project']}", "", f"**Deliverable:** {w['deliverable']}", ""]
        lines += ["### Knowledge checkpoint (close-book)", ""] + [f"- [ ] {c}" for c in w["checkpoint"]] + [""]
        lines += ["### Day by day", "", "| Date | Phase | Type | Task | What to do | Link | h | Output |", "|---|---|---|---|---|---|---|---|"]
        for i, day in enumerate(w["days"]):
            d = s + timedelta(days=i)
            for t in day:
                phase, typ, topic, desc, link, hours, deliv = t
                link_md = md_link("link", link) if link else ""
                lines.append(f"| {fmt(d)} | {phase} | {typ} | **{esc(topic)}** | {esc(desc)} | {link_md} | {hours:g} | {esc(deliv)} |")
        lines.append("")
    return "\n".join(lines) + "\n"


def render_papers_md():
    level_names = {1: "Level 1 - accessible (read in one sitting)", 2: "Level 2 - core (some notation)",
                   3: "Level 3 - technical (derivations / systems detail)", 4: "Level 4 - advanced (math- or systems-heavy)",
                   5: "Level 5 - frontier"}
    core = [p for p in PAPERS if p[10] == "CORE"]
    ext = [p for p in PAPERS if p[10] == "EXT"]
    lines = [
        "# 7. Paper reading list (easiest -> hardest)",
        "",
        "> Generated by `tools/build.py` from `tools/plan_data.py`.",
        "",
        f"{len(core)} core papers are scheduled in 2026 (the 'Week' column says when); {len(ext)} extended papers move to 2027. "
        "Within the schedule you read them in week order - this page orders them by difficulty so you can see the climb.",
        "",
        "## How to read every paper (the protocol)",
        "",
        "1. **Pass 1 (10 min):** title, abstract, intro, section headings, figures, conclusion. Decide: what problem, what claim, what evidence.",
        "2. **Pass 2 (45-60 min):** read the sections listed below carefully; skip proofs; write down every tensor shape / equation you would need to implement it.",
        "3. **Pass 3 (only for papers with an 'Implement' entry):** re-derive the key equation on paper, implement the core idea in <200 lines, run the smallest experiment that could falsify your understanding.",
        "4. **Note template** (`papers/NN-name.md`): problem - key idea in 3 sentences - the one equation - what I implemented - what surprised me - open questions.",
        "",
        "Reproductions scheduled in 2026 (small-scale, 1 GPU): #1 architecture ablations (W3), #2 mini IsoFLOP / Chinchilla (W4), "
        "#3 BEIR BM25 baseline (W5), #4 LoRA learning-rate finding (W8), #5 Dr. GRPO length bias (W9), #6 continuous batching gains (W11).",
        "",
    ]
    for lvl in sorted({p[4] for p in core}):
        group = sorted([p for p in core if p[4] == lvl], key=lambda p: (int(p[5][1:]) if p[5].startswith("W") else 99))
        lines += [f"## {level_names[lvl]}", ""]
        for p in group:
            key, title, authors, year, diff, week, why, focus, impl, repro, _ = p
            url, _ = resolve(key)
            lines += [
                f"### {title}",
                f"*{authors}, {year}* - [{url}]({url}) - difficulty {diff}/5 - **{week}**",
                "",
                f"- **Why it matters:** {why}",
                f"- **Prerequisites:** {PAPER_PREREQS[key]}",
                f"- **Read carefully:** {focus}",
                f"- **Implement:** {impl}",
                f"- **Reproduce:** {repro}",
                "",
            ]
    lines += ["## Extended list (2027)", "", "| Paper | Authors (year) | Difficulty | Why | Focus | Prerequisites |", "|---|---|---|---|---|---|"]
    for p in sorted(ext, key=lambda p: p[4]):
        url, _ = resolve(p[0])
        lines.append(f"| [{esc(p[1])}]({url}) | {esc(p[2])} ({p[3]}) | {p[4]}/5 | {esc(p[6])} | {esc(p[7])} | {esc(PAPER_PREREQS[p[0]])} |")
    return "\n".join(lines) + "\n"


# ---------------------------------------------------------------- excel
FONT = "Arial"
HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
TITLE_FONT = Font(name=FONT, bold=True, size=16, color="1F3864")
H2_FONT = Font(name=FONT, bold=True, size=12, color="1F3864")
BODY = Font(name=FONT, size=10)
BOLD = Font(name=FONT, size=10, bold=True)
LINK_FONT = Font(name=FONT, size=10, color="0563C1", underline="single")
INPUT_FONT = Font(name=FONT, size=11, color="0000FF", bold=True)
BAND = PatternFill("solid", fgColor="F2F5FA")
DONE_FILL = PatternFill("solid", fgColor="D9EAD3")
YELLOW = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="D0D7E5")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)
UNCHECKED, CHECKED = "☐", "☑"


def header(ws, row, labels, widths=None):
    for c, lbl in enumerate(labels, 1):
        cell = ws.cell(row=row, column=c, value=lbl)
        cell.font, cell.fill, cell.alignment, cell.border = HDR_FONT, HDR_FILL, CENTER, BORDER
    if widths:
        for c, wdt in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(c)].width = wdt


def done_count(rng):
    """Count ticked cells whether the user uses the dropdown or Excel's native checkbox (TRUE)."""
    return f'(COUNTIF({rng},"{CHECKED}")+COUNTIF({rng},TRUE))'


def build_xlsx(path):
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    plan = wb.create_sheet("Daily Plan")
    projects = wb.create_sheet("Projects")
    papers = wb.create_sheet("Papers")
    resources = wb.create_sheet("Resources")

    # ---------------- Daily Plan
    data = rows()
    cols = ["#", "Date", "Day", "Week", "Phase", "Type", "Topic", "Description", "Link", "Hours",
            "Deliverable / assignment", "Done", "Notes"]
    header(plan, 1, cols, [5, 11, 6, 7, 7, 11, 30, 70, 32, 7, 26, 7, 30])
    dv = DataValidation(type="list", formula1=f'"{UNCHECKED},{CHECKED}"', allow_blank=False)
    dv.error, dv.errorTitle = f"Pick {UNCHECKED} or {CHECKED}", "Done"
    plan.add_data_validation(dv)
    for i, r in enumerate(data, start=2):
        url, _ = resolve(r["link"])
        vals = [i - 1, r["date"], r["day"], f"W{r['week']}", r["phase"], r["type"], r["topic"], r["desc"],
                url, r["hours"], r["deliv"], UNCHECKED, ""]
        for c, v in enumerate(vals, 1):
            cell = plan.cell(row=i, column=c, value=v)
            cell.font, cell.alignment, cell.border = BODY, WRAP, BORDER
            if r["week"] % 2 == 0:
                cell.fill = BAND
        plan.cell(row=i, column=2).number_format = "ddd dd mmm yyyy"
        plan.cell(row=i, column=7).font = BOLD
        if url:
            lc = plan.cell(row=i, column=9)
            lc.hyperlink, lc.font = url, LINK_FONT
        dc = plan.cell(row=i, column=12)
        dc.alignment, dc.font = Alignment(horizontal="center", vertical="top"), Font(name=FONT, size=14)
        dv.add(dc)
    last = len(data) + 1
    plan.freeze_panes = "H2"
    plan.auto_filter.ref = f"A1:M{last}"
    plan.conditional_formatting.add(
        f"A2:M{last}", FormulaRule(formula=[f'OR($L2="{CHECKED}",$L2=TRUE)'], fill=DONE_FILL, stopIfTrue=True))
    plan.conditional_formatting.add(
        f"B2:B{last}", FormulaRule(formula=[f'AND($B2<TODAY(),$L2="{UNCHECKED}")'], font=Font(name=FONT, color="C00000", bold=True)))

    rng_done = f"'Daily Plan'!$L$2:$L${last}"
    rng_hours = f"'Daily Plan'!$J$2:$J${last}"
    rng_phase = f"'Daily Plan'!$E$2:$E${last}"
    rng_week = f"'Daily Plan'!$D$2:$D${last}"
    rng_date = f"'Daily Plan'!$B$2:$B${last}"

    # ---------------- Projects
    header(projects, 1, ["ID", "Project", "Weeks", "Due", "Deliverable", "Skills developed", "Status", "Done", "Your repo link / notes"],
           [6, 34, 9, 12, 60, 50, 13, 7, 36])
    pdv = DataValidation(type="list", formula1=f'"{UNCHECKED},{CHECKED}"')
    sdv = DataValidation(type="list", formula1='"Not started,In progress,Blocked,Done"')
    projects.add_data_validation(pdv)
    projects.add_data_validation(sdv)
    from datetime import date as _date
    for i, (pid, name, weeks, due, deliv, skills) in enumerate(PROJECTS, start=2):
        vals = [pid, name, weeks, _date.fromisoformat(due), deliv, skills, "Not started", UNCHECKED, ""]
        for c, v in enumerate(vals, 1):
            cell = projects.cell(row=i, column=c, value=v)
            cell.font, cell.alignment, cell.border = BODY, WRAP, BORDER
        projects.cell(row=i, column=2).font = BOLD
        projects.cell(row=i, column=4).number_format = "dd mmm yyyy"
        projects.cell(row=i, column=8).alignment = Alignment(horizontal="center", vertical="top")
        projects.cell(row=i, column=8).font = Font(name=FONT, size=14)
        sdv.add(projects.cell(row=i, column=7))
        pdv.add(projects.cell(row=i, column=8))
    plast = len(PROJECTS) + 1
    projects.freeze_panes = "C2"
    projects.conditional_formatting.add(
        f"A2:I{plast}", FormulaRule(formula=[f'OR($H2="{CHECKED}",$H2=TRUE,$G2="Done")'], fill=DONE_FILL))

    # ---------------- Papers
    header(papers, 1, ["#", "Title", "Authors", "Year", "Difficulty (1-5)", "Week", "Link", "Why it matters",
                       "Prerequisites", "Read carefully", "Implement", "Reproduce", "Tier", "Read", "Notes"],
           [5, 40, 28, 6, 9, 7, 32, 40, 30, 32, 28, 28, 7, 7, 30])
    rdv = DataValidation(type="list", formula1=f'"{UNCHECKED},{CHECKED}"')
    papers.add_data_validation(rdv)
    ordered = sorted(PAPERS, key=lambda p: (p[10] != "CORE", p[4], int(p[5][1:]) if p[5].startswith("W") else 99))
    for i, p in enumerate(ordered, start=2):
        key, title, authors, year, diff, week, why, focus, impl, repro, tier = p
        url, _ = resolve(key)
        vals = [i - 1, title, authors, year, diff, week, url, why, PAPER_PREREQS[key], focus, impl, repro, tier, UNCHECKED, ""]
        for c, v in enumerate(vals, 1):
            cell = papers.cell(row=i, column=c, value=v)
            cell.font, cell.alignment, cell.border = BODY, WRAP, BORDER
        papers.cell(row=i, column=2).font = BOLD
        lc = papers.cell(row=i, column=7)
        lc.hyperlink, lc.font = url, LINK_FONT
        papers.cell(row=i, column=14).alignment = Alignment(horizontal="center", vertical="top")
        papers.cell(row=i, column=14).font = Font(name=FONT, size=14)
        rdv.add(papers.cell(row=i, column=14))
    palast = len(PAPERS) + 1
    papers.freeze_panes = "C2"
    papers.auto_filter.ref = f"A1:O{palast}"
    papers.conditional_formatting.add(
        f"A2:O{palast}", FormulaRule(formula=[f'OR($N2="{CHECKED}",$N2=TRUE)'], fill=DONE_FILL))

    # ---------------- Resources
    header(resources, 1, ["Kind", "Title", "By / channel", "Duration", "Link", "Week(s)", "Why / notes"],
           [11, 52, 30, 11, 50, 12, 60])
    res = OrderedDict()
    for w in WEEKS:
        for kind in ("primary", "secondary", "reference"):
            for label, link in w[kind]:
                url, _ = resolve(link)
                entry = res.setdefault(url, dict(kind=kind.upper(), title=label, by="", dur="", weeks=[], why=""))
                entry["weeks"].append(f"W{w['n']}")
        for title, chan, link, dur, why, sec in w["videos"]:
            url, _ = resolve(link)
            entry = res.setdefault(url, dict(kind="VIDEO", title=title, by=chan, dur=dur, weeks=[], why=f"{why} ({sec})"))
            entry["kind"], entry["by"], entry["dur"], entry["why"] = "VIDEO", chan, dur, f"{why} ({sec})"
            entry["weeks"].append(f"W{w['n']}")
    for i, (url, e) in enumerate(res.items(), start=2):
        weeks = ", ".join(OrderedDict.fromkeys(e["weeks"]))
        vals = [e["kind"], e["title"], e["by"], e["dur"], url, weeks, e["why"]]
        for c, v in enumerate(vals, 1):
            cell = resources.cell(row=i, column=c, value=v)
            cell.font, cell.alignment, cell.border = BODY, WRAP, BORDER
        lc = resources.cell(row=i, column=5)
        lc.hyperlink, lc.font = url, LINK_FONT
    resources.freeze_panes = "C2"
    resources.auto_filter.ref = f"A1:G{len(res) + 1}"

    # ---------------- Summary
    ws = summary
    for col, wdt in zip("ABCDEFGHI", [44, 16, 14, 14, 14, 14, 14, 3, 3]):
        ws.column_dimensions[col].width = wdt
    ws["A1"] = "AI Software Engineer - 2026 sprint plan"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = (f"{START:%d %b %Y} -> {END:%d %b %Y}. Tick tasks in 'Daily Plan' column L (choose {CHECKED}). "
                "Every number on this page is a formula over the other sheets.")
    ws["A2"].font = Font(name=FONT, size=10, italic=True)
    ws["A2"].alignment = WRAP
    ws.merge_cells("A2:G2")
    ws.row_dimensions[2].height = 28
    ws["A3"] = (f"How to use: only edit the blue/yellow cells - Done ({UNCHECKED}/{CHECKED}) in Daily Plan / Projects / Papers, "
                "Status in Projects, and Notes. Excel 365 users may convert the Done column to native checkboxes "
                "(Insert > Checkbox); the formulas count both.")
    ws["A3"].font = Font(name=FONT, size=9, color="595959")
    ws["A3"].alignment = WRAP
    ws.merge_cells("A3:G3")
    ws.row_dimensions[3].height = 30

    r = 5
    ws.cell(row=r, column=1, value="Overall progress").font = H2_FONT
    r += 1
    kpis = [
        ("Start date", f"=MIN({rng_date})", "dd mmm yyyy"),
        ("End date", f"=MAX({rng_date})", "dd mmm yyyy"),
        ("Study days", f"=SUMPRODUCT(1/COUNTIF({rng_date},{rng_date}))", "0"),
        ("Tasks planned", f"=COUNTA('Daily Plan'!$G$2:$G${last})", "0"),
        ("Tasks done", f"={done_count(rng_done)}", "0"),
        ("Tasks done (%)", f"=IFERROR(B{r + 4}/B{r + 3},0)", "0.0%"),
        ("Hours planned", f"=SUM({rng_hours})", "0.0"),
        ("Hours done", f'=SUMIF({rng_done},"{CHECKED}",{rng_hours})+SUMIF({rng_done},TRUE,{rng_hours})', "0.0"),
        ("Hours done (%)", f"=IFERROR(B{r + 7}/B{r + 6},0)", "0.0%"),
        ("Tasks scheduled up to today", f'=COUNTIF({rng_date},"<="&TODAY())', "0"),
        ("Ahead (+) / behind (-) in tasks", f"=B{r + 4}-B{r + 9}", "+0;-0;0"),
        ("Projects done", f'={done_count("Projects!$H$2:$H$" + str(plast))}', "0"),
        ("Papers read", f'={done_count("Papers!$N$2:$N$" + str(palast))}', "0"),
    ]
    for label, formula, numfmt in kpis:
        a = ws.cell(row=r, column=1, value=label)
        b = ws.cell(row=r, column=2, value=formula)
        a.font, b.font, b.number_format = BODY, BOLD, numfmt
        a.border = b.border = BORDER
        r += 1
    ws.cell(row=r, column=1, value=f"Assumptions: {WEEKDAY_HOURS} h Mon-Fri, {SATURDAY_HOURS} h Sat, Sunday rest; "
                                   "25 Dec is a 0-hour buffer. Source: plan design (tools/plan_data.py).").font = Font(name=FONT, size=9, italic=True, color="595959")
    r += 2

    # progress by phase
    ws.cell(row=r, column=1, value="Progress by phase").font = H2_FONT
    r += 1
    for c, lbl in enumerate(["Phase", "Code", "Tasks", "Done", "Done %", "Hours", "Hours done"], 1):
        cell = ws.cell(row=r, column=c, value=lbl)
        cell.font, cell.fill, cell.alignment, cell.border = HDR_FONT, HDR_FILL, CENTER, BORDER
    r += 1
    for code, name in PHASES.items():
        vals = [name, code,
                f'=COUNTIF({rng_phase},B{r})',
                f'=COUNTIFS({rng_phase},B{r},{rng_done},"{CHECKED}")+COUNTIFS({rng_phase},B{r},{rng_done},TRUE)',
                f'=IFERROR(D{r}/C{r},0)',
                f'=SUMIF({rng_phase},B{r},{rng_hours})',
                f'=SUMIFS({rng_hours},{rng_phase},B{r},{rng_done},"{CHECKED}")+SUMIFS({rng_hours},{rng_phase},B{r},{rng_done},TRUE)']
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font, cell.border = BODY, BORDER
        ws.cell(row=r, column=5).number_format = "0%"
        ws.cell(row=r, column=6).number_format = "0.0"
        ws.cell(row=r, column=7).number_format = "0.0"
        r += 1
    r += 1

    # progress by week
    ws.cell(row=r, column=1, value="Progress by week").font = H2_FONT
    r += 1
    for c, lbl in enumerate(["Week focus", "Week", "Starts", "Tasks", "Done", "Done %", "Hours"], 1):
        cell = ws.cell(row=r, column=c, value=lbl)
        cell.font, cell.fill, cell.alignment, cell.border = HDR_FONT, HDR_FILL, CENTER, BORDER
    r += 1
    for w in WEEKS:
        s, _ = week_dates(w)
        vals = [w["title"], f"W{w['n']}", s,
                f"=COUNTIF({rng_week},B{r})",
                f'=COUNTIFS({rng_week},B{r},{rng_done},"{CHECKED}")+COUNTIFS({rng_week},B{r},{rng_done},TRUE)',
                f"=IFERROR(E{r}/D{r},0)",
                f"=SUMIF({rng_week},B{r},{rng_hours})"]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font, cell.border, cell.alignment = BODY, BORDER, WRAP
        ws.cell(row=r, column=3).number_format = "dd mmm"
        ws.cell(row=r, column=6).number_format = "0%"
        ws.cell(row=r, column=7).number_format = "0.0"
        r += 1
    r += 1

    # skills: already known
    ws.cell(row=r, column=1, value="What you already know (skip or quick revision only)").font = H2_FONT
    r += 1
    for c, lbl in enumerate(["Topic", "Action in this plan"], 1):
        cell = ws.cell(row=r, column=c, value=lbl)
        cell.font, cell.fill, cell.border = HDR_FONT, HDR_FILL, BORDER
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
    r += 1
    for topic, action in KNOWN:
        ws.cell(row=r, column=1, value=topic).font = BODY
        ws.cell(row=r, column=2, value=action).font = BODY
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
        for c in (1, 2):
            ws.cell(row=r, column=c).border, ws.cell(row=r, column=c).alignment = BORDER, WRAP
        r += 1
    r += 1

    # skills: learning
    ws.cell(row=r, column=1, value="What you are learning (target level) and deferred items").font = H2_FONT
    r += 1
    ws.cell(row=r, column=1, value="MASTER = implement + explain without notes; WORKING = explain + use a library + one experiment; "
                                   "DEFER-2027 = scheduled in the 2027 continuation plan.").font = Font(name=FONT, size=9, italic=True, color="595959")
    r += 1
    for c, lbl in enumerate(["Topic", "Phase", "Target level", "When", "Self-rating (1-5)"], 1):
        cell = ws.cell(row=r, column=c, value=lbl)
        cell.font, cell.fill, cell.alignment, cell.border = HDR_FONT, HDR_FILL, CENTER, BORDER
    r += 1
    rating = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True)
    ws.add_data_validation(rating)
    level_fill = {"MASTER": PatternFill("solid", fgColor="D9EAD3"), "WORKING": PatternFill("solid", fgColor="DEEAF6"),
                  "DEFER-2027": PatternFill("solid", fgColor="EDEDED")}
    for topic, phase, level, when in LEARN:
        vals = [topic, phase, level, when, None]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font, cell.border, cell.alignment = BODY, BORDER, WRAP
        ws.cell(row=r, column=3).fill = level_fill[level]
        rc = ws.cell(row=r, column=5)
        rc.fill, rc.font = YELLOW, INPUT_FONT
        rating.add(rc)
        r += 1
    ws.cell(row=r + 1, column=1, value="Example: rate 'KV cache, prefill vs decode' as 4 once you can derive the memory formula close-book.").font = Font(name=FONT, size=9, italic=True, color="595959")

    # mark editable cells on the other sheets
    for i in range(2, last + 1):
        plan.cell(row=i, column=12).font = Font(name=FONT, size=14, color="0000FF")
    for i in range(2, plast + 1):
        projects.cell(row=i, column=7).fill = YELLOW
    for sheet in wb.worksheets:
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        if sheet is not ws:
            sheet.print_title_rows = "1:1"
    wb.save(path)
    return len(data)


def main():
    problems = validate()
    if problems:
        print("VALIDATION PROBLEMS:")
        for p in problems:
            print(" -", p)
        sys.exit(1)
    os.chdir(ROOT)
    with open("curriculum/04-week-by-week.md", "w") as f:
        f.write(render_weeks_md())
    with open("curriculum/07-paper-reading-list.md", "w") as f:
        f.write(render_papers_md())
    n = build_xlsx("tracker/AI_Engineer_2026_Plan.xlsx")
    all_rows = rows()
    print(f"weeks={len(WEEKS)} tasks={n} days={len({r['date'] for r in all_rows})} hours={sum(r['hours'] for r in all_rows)}")
    by_phase = OrderedDict((k, 0.0) for k in PHASES)
    for r in all_rows:
        by_phase[r["phase"]] += r["hours"]
    for k, v in by_phase.items():
        print(f"  {k:4} {v:5.1f} h")


if __name__ == "__main__":
    main()
